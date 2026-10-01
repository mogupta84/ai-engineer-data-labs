"""
FreshCart RAG labs: shared library.

Runs offline by default so a cohort is never blocked by API keys or quota.
Set FRESHCART_EMBEDDER=openai|azure to use a hosted model instead.
"""
from __future__ import annotations
import os, re, json, hashlib, pathlib
from dataclasses import dataclass, field, asdict
from typing import Iterable

import numpy as np

ROOT = pathlib.Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"


# ----------------------------------------------------------------------------
# Documents and chunks
# ----------------------------------------------------------------------------
@dataclass
class Doc:
    text: str
    meta: dict = field(default_factory=dict)


@dataclass
class Chunk:
    id: str
    text: str
    meta: dict = field(default_factory=dict)

    def to_json(self):
        return {"id": self.id, "text": self.text, "meta": self.meta}


# ----------------------------------------------------------------------------
# Embedders
# ----------------------------------------------------------------------------
class TfidfSvdEmbedder:
    """Offline embedder. Not state of the art; entirely adequate for the lab
    mechanics, and it means nobody waits on an API key."""

    name = "tfidf-svd-256"

    def __init__(self, dim: int = 256):
        from sklearn.feature_extraction.text import TfidfVectorizer
        from sklearn.decomposition import TruncatedSVD
        self.dim = dim
        self.name = f"tfidf-svd-{dim}"
        self.vec = TfidfVectorizer(lowercase=True, stop_words="english",
                                   ngram_range=(1, 2), min_df=1)
        self.svd = TruncatedSVD(n_components=dim, random_state=0)
        self.fitted = False

    def fit(self, texts: list[str]):
        X = self.vec.fit_transform(texts)
        n = min(self.dim, max(2, min(X.shape) - 1))
        if n != self.dim:
            from sklearn.decomposition import TruncatedSVD
            self.svd = TruncatedSVD(n_components=n, random_state=0)
            self.dim = n
        # A small corpus allows fewer dimensions than asked for (30 pieces give at most 29),
        # so the name carries the number actually used.
        self.name = f"tfidf-svd-{self.dim}"
        self.svd.fit(X)
        self.fitted = True
        return self

    def encode(self, texts: list[str]) -> np.ndarray:
        if not self.fitted:
            raise RuntimeError("call fit() on the corpus before encode()")
        X = self.vec.transform(texts)
        V = self.svd.transform(X).astype("float32")
        norms = np.linalg.norm(V, axis=1, keepdims=True)
        norms[norms == 0] = 1.0
        return V / norms


class HashingEmbedder:
    """Offline embedder that needs no fitting, so a chunk's vector never changes
    when OTHER documents change. Lab 9 uses it to prove vector reuse offline.
    (TF-IDF + SVD must be re-fitted on the whole corpus after any change, which
    silently re-embeds everything: a lesson of its own.)"""

    name = "hashing-256"

    def __init__(self, dim: int = 256, features: int = 2 ** 14):
        from sklearn.feature_extraction.text import HashingVectorizer
        self.dim = dim
        self.vec = HashingVectorizer(n_features=features, alternate_sign=False,
                                     ngram_range=(1, 2), stop_words="english", norm="l2")
        rng = np.random.default_rng(42)          # fixed seed: same text, same vector, every run
        self.proj = (rng.standard_normal((features, dim)) / np.sqrt(dim)).astype("float32")

    def fit(self, texts):
        return self

    def encode(self, texts: list[str]) -> np.ndarray:
        X = self.vec.transform(texts)
        V = (X @ self.proj).astype("float32")
        norms = np.linalg.norm(V, axis=1, keepdims=True)
        norms[norms == 0] = 1.0
        return V / norms


class FastEmbedEmbedder:
    """A current open-source retrieval model, BGE small, through fastembed.
    No key and no account: the ONNX model downloads once per session, about
    130 MB, and every embedding after that is computed locally. Needs no
    fitting, because it was trained on far more text than this corpus."""

    name = "bge-small-en-v1.5"

    def __init__(self, model: str = "BAAI/bge-small-en-v1.5"):
        # Keep the download quiet. Without these the model download prints two progress
        # bars, a byte count and a notice about setting HF_TOKEN, in the middle of the
        # comparison output. None of it is an error and none of it needs a token. These
        # must be set before huggingface_hub is imported, which is why they are here.
        import os, warnings
        os.environ.setdefault("HF_HUB_DISABLE_PROGRESS_BARS", "1")
        os.environ.setdefault("HF_HUB_DISABLE_TELEMETRY", "1")
        os.environ.setdefault("HF_HUB_DISABLE_XET", "1")
        warnings.filterwarnings("ignore", message="Cannot enable progress bars")
        from fastembed import TextEmbedding
        self.model = TextEmbedding(model)
        self.name = model.split("/")[-1]

    def fit(self, texts):
        return self

    def encode(self, texts: list[str]) -> np.ndarray:
        V = np.array(list(self.model.embed(list(texts))), dtype="float32")
        norms = np.linalg.norm(V, axis=1, keepdims=True)
        norms[norms == 0] = 1.0
        return V / norms


class OpenAIEmbedder:
    """Hosted embeddings. With azure=True it calls your Azure OpenAI deployment
    through the v1 endpoint, using the settings in lab.env.

    Vectors are cached on disk by text hash, so re-running a lab does not pay to
    embed the same chunk twice. That is the Lab 9 idea (reuse unchanged vectors)
    applied from day one."""
    name = "text-embedding-3-small"

    def __init__(self, model: str | None = None, azure: bool = False):
        from openai import OpenAI
        self.azure = azure
        if azure:
            self.client = OpenAI(base_url=os.environ["AZURE_OPENAI_BASE_URL"],
                                 api_key=os.environ["AZURE_OPENAI_API_KEY"], max_retries=5)
            self.model = model or os.getenv("AZURE_OPENAI_EMBED_DEPLOYMENT", self.name)
        else:
            self.client = OpenAI()
            self.model = model or os.getenv("FRESHCART_EMBED_MODEL", self.name)
        self.name = self.model
        self.tokens = 0          # tokens actually sent (cache hits cost nothing)
        self.cache_hits = 0
        self._count = __import__("threading").Lock()
        self._cache = _EmbeddingCache(ROOT / "data" / "cache" / "embeddings.sqlite", self.model)

    def fit(self, texts): return self

    def encode(self, texts: list[str]) -> np.ndarray:
        found = self._cache.get_many(texts)
        todo = [t for t in dict.fromkeys(texts) if t not in found]
        with self._count:
            self.cache_hits += len(texts) - len(todo)
        for i in range(0, len(todo), 64):
            batch = todo[i:i + 64]
            r = self.client.embeddings.create(model=self.model, input=batch)
            used = getattr(r.usage, "prompt_tokens", 0) or 0
            with self._count:
                self.tokens += used
            if self.azure:                       # so every lab's cost line includes embeddings
                from azure_backends import USAGE
                with USAGE.lock:
                    USAGE.embed += used
            fresh = {t: d.embedding for t, d in zip(batch, r.data)}
            self._cache.put_many(fresh)
            found.update(fresh)
        V = np.array([found[t] for t in texts], dtype="float32")
        return V / np.linalg.norm(V, axis=1, keepdims=True)


class _EmbeddingCache:
    def __init__(self, path: pathlib.Path, model: str):
        import sqlite3, threading
        path.parent.mkdir(parents=True, exist_ok=True)
        self.model = model
        self.lock = threading.Lock()
        self.con = sqlite3.connect(str(path), check_same_thread=False, timeout=30)
        self.con.execute("CREATE TABLE IF NOT EXISTS emb (key TEXT PRIMARY KEY, v BLOB)")

    def _key(self, text):
        return hashlib.sha1(f"{self.model}|{text}".encode("utf-8")).hexdigest()

    def get_many(self, texts):
        out = {}
        with self.lock:
            for t in dict.fromkeys(texts):
                row = self.con.execute("SELECT v FROM emb WHERE key=?", (self._key(t),)).fetchone()
                if row:
                    out[t] = np.frombuffer(row[0], dtype="float32").tolist()
        return out

    def put_many(self, mapping):
        with self.lock:
            self.con.executemany("INSERT OR REPLACE INTO emb VALUES (?,?)",
                                 [(self._key(t), np.array(v, dtype="float32").tobytes())
                                  for t, v in mapping.items()])
            self.con.commit()


def get_embedder():
    kind = os.getenv("FRESHCART_EMBEDDER", "offline").lower()
    if kind == "azure":
        return OpenAIEmbedder(azure=True)
    if kind == "openai":
        return OpenAIEmbedder()
    if kind == "hashing":
        return HashingEmbedder()
    if kind == "fastembed":
        return FastEmbedEmbedder()
    return TfidfSvdEmbedder()


# ----------------------------------------------------------------------------
# Loaders: one per format
# ----------------------------------------------------------------------------
def load_pdf(path: pathlib.Path) -> list[Doc]:
    from pypdf import PdfReader
    docs = []
    for i, page in enumerate(PdfReader(str(path)).pages, start=1):
        docs.append(Doc(text=page.extract_text() or "",
                        meta={"source": path.name, "page": i, "format": "pdf"}))
    return docs


def load_docx(path: pathlib.Path) -> list[Doc]:
    """Tables are rendered row-wise with headers repeated, so a retrieved row
    still carries its column names. Naive extraction loses them."""
    from docx import Document as Docx
    d = Docx(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t_i, table in enumerate(d.tables):
        rows = [[c.text.strip() for c in r.cells] for r in table.rows]
        if not rows:
            continue
        header, body = rows[0], rows[1:]
        for r in body:
            pairs = "; ".join(f"{h}: {v}" for h, v in zip(header, r) if v)
            parts.append(f"[table {t_i + 1}] {pairs}")
    return [Doc(text="\n\n".join(parts),
                meta={"source": path.name, "format": "docx"})]


def load_pptx(path: pathlib.Path) -> list[Doc]:
    """One Doc per slide. Shapes are sorted by vertical then horizontal position
    so reading order survives, which creation order does not guarantee."""
    from pptx import Presentation
    prs = Presentation(str(path))
    docs = []
    for i, slide in enumerate(prs.slides, start=1):
        shapes = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        shapes.sort(key=lambda sh: (sh.top or 0, sh.left or 0))
        text = "\n".join(sh.text_frame.text.strip() for sh in shapes)
        docs.append(Doc(text=text, meta={"source": path.name, "slide": i, "format": "pptx"}))
    return docs


def load_xlsx(path: pathlib.Path) -> list[Doc]:
    """Finds the real header row rather than assuming row 1, then emits one Doc
    per data row with headers repeated."""
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    docs = []
    for ws in wb.worksheets:
        grid = [[c.value for c in row] for row in ws.iter_rows()]
        header_idx, header = None, None
        for i, row in enumerate(grid[:15]):
            filled = [c for c in row if c not in (None, "")]
            if len(filled) >= 3 and all(isinstance(c, str) for c in filled):
                header_idx, header = i, [str(c).strip() if c else "" for c in row]
                break
        if header_idx is None:
            continue
        for r in grid[header_idx + 1:]:
            if not any(c not in (None, "") for c in r):
                continue
            pairs = "; ".join(f"{h}: {v}" for h, v in zip(header, r) if h and v not in (None, ""))
            if pairs:
                docs.append(Doc(text=pairs, meta={"source": path.name, "sheet": ws.title,
                                                  "format": "xlsx"}))
    return docs


LOADERS = {".pdf": load_pdf, ".docx": load_docx, ".pptx": load_pptx, ".xlsx": load_xlsx}


def load_all(raw_dir: pathlib.Path = RAW) -> list[Doc]:
    docs = []
    for p in sorted(raw_dir.iterdir()):
        fn = LOADERS.get(p.suffix.lower())
        if fn:
            docs.extend(fn(p))
    return docs


# ----------------------------------------------------------------------------
# Extraction verification: the step tutorials skip
# ----------------------------------------------------------------------------
def verify(docs: list[Doc]) -> list[dict]:
    """Returns one report row per source. Empty text is the failure that matters:
    it raises no exception anywhere in the pipeline."""
    by_source: dict[str, list[Doc]] = {}
    for d in docs:
        by_source.setdefault(d.meta.get("source", "?"), []).append(d)
    report = []
    for src, ds in sorted(by_source.items()):
        chars = sum(len(d.text.strip()) for d in ds)
        empty = sum(1 for d in ds if not d.text.strip())
        report.append({
            "source": src, "units": len(ds), "chars": chars, "empty_units": empty,
            "status": "EMPTY, needs OCR" if chars == 0 else
                      ("thin" if chars < 400 else "ok"),
        })
    return report


# ----------------------------------------------------------------------------
# Chunking
# ----------------------------------------------------------------------------
def chunk_fixed(text: str, size: int = 500, overlap: int = 0) -> list[str]:
    out, i = [], 0
    while i < len(text):
        out.append(text[i:i + size])
        i += max(1, size - overlap)
    return [c for c in out if c.strip()]


def chunk_recursive(text: str, size: int = 900, overlap: int = 120) -> list[str]:
    """Prefer paragraph boundaries, then sentences, then characters."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks, buf = [], ""
    for p in paras:
        if len(buf) + len(p) + 2 <= size:
            buf = f"{buf}\n\n{p}".strip()
            continue
        if buf:
            chunks.append(buf)
        if len(p) <= size:
            buf = p
        else:
            sents = re.split(r"(?<=[.!?])\s+", p)
            buf = ""
            for s in sents:
                if len(buf) + len(s) + 1 <= size:
                    buf = f"{buf} {s}".strip()
                else:
                    if buf:
                        chunks.append(buf)
                    buf = s if len(s) <= size else ""
                    if len(s) > size:
                        chunks.extend(chunk_fixed(s, size, overlap))
    if buf:
        chunks.append(buf)
    if overlap > 0 and len(chunks) > 1:
        out = [chunks[0]]
        for prev, cur in zip(chunks, chunks[1:]):
            out.append((prev[-overlap:] + " " + cur).strip())
        chunks = out
    return chunks


HEADING = re.compile(r"^\s*(\d+(?:\.\d+)*)\.?\s+([A-Z][^\n]{3,80})$", re.M)


def chunk_structural(text: str, size: int = 1200) -> list[str]:
    """Split on numbered headings and prepend the heading to each chunk so the
    chunk still says what it is about when retrieved on its own."""
    marks = [(m.start(), m.group(0).strip()) for m in HEADING.finditer(text)]
    if not marks:
        return chunk_recursive(text, size=size)
    chunks = []
    if marks[0][0] > 0:
        pre = text[:marks[0][0]].strip()
        if pre:
            chunks.extend(chunk_recursive(pre, size=size))
    for (start, head), nxt in zip(marks, marks[1:] + [(len(text), "")]):
        body = text[start:nxt[0]].strip()
        if len(body) <= size:
            chunks.append(body)
        else:
            for piece in chunk_recursive(body, size=size):
                chunks.append(piece if piece.startswith(head) else f"{head}\n{piece}")
    return [c for c in chunks if c.strip()]


STRATEGIES = {"fixed": chunk_fixed, "recursive": chunk_recursive, "structural": chunk_structural}


def strategy_for(doc: Doc) -> str:
    """Per document type, not one rule for everything."""
    fmt = doc.meta.get("format")
    if fmt == "xlsx":
        return "row"          # already one row per Doc
    if fmt == "pptx":
        return "slide"        # already one slide per Doc
    if fmt == "pdf" and "policy" in doc.meta.get("source", ""):
        return "structural"
    return "recursive"


def chunk_docs(docs: list[Doc], override: str | None = None) -> list[Chunk]:
    chunks: list[Chunk] = []
    for d in docs:
        if not d.text.strip():
            continue
        strat = override or strategy_for(d)
        pieces = [d.text] if strat in ("row", "slide") else STRATEGIES[strat](d.text)
        for i, piece in enumerate(pieces):
            title = d.meta.get("source", "")
            prefix = ""
            if d.meta.get("slide"):
                prefix = f"[{title} slide {d.meta['slide']}] "
            elif d.meta.get("page"):
                prefix = f"[{title} p{d.meta['page']}] "
            body = prefix + piece
            cid = hashlib.sha1(f"{title}|{d.meta.get('page') or d.meta.get('slide')}|{i}|{piece[:60]}"
                               .encode()).hexdigest()[:12]
            meta = dict(d.meta)
            meta.update({"strategy": strat, "chunk_index": i})
            meta.update(enrich(title))
            chunks.append(Chunk(id=cid, text=body, meta=meta))
    return chunks


def enrich(source: str) -> dict:
    """Metadata attached at ingestion. Retrofitting this later means reprocessing
    the corpus, which is why it belongs here."""
    m = {}
    if "refund_policy_v3_2024" in source:
        m.update(doc_type="policy", effective_date="2024-04-01", version="3.0", superseded=False)
    elif "refund_policy_v2_2022" in source:
        m.update(doc_type="policy", effective_date="2022-06-15", version="2.0", superseded=True)
    elif "partner_agreement" in source:
        m.update(doc_type="contract", effective_date="2023-01-01", superseded=False)
    elif "sla_matrix" in source:
        m.update(doc_type="sla", effective_date="2024-04-01", superseded=False)
    elif "escalation_sop" in source:
        m.update(doc_type="sop", effective_date="2024-03-01", superseded=False)
    elif "qbr" in source:
        m.update(doc_type="report", effective_date="2024-04-12", superseded=False)
    elif "refund_claims" in source:
        m.update(doc_type="register", effective_date="2024-04-01", superseded=False)
    return m


# ----------------------------------------------------------------------------
# Index
# ----------------------------------------------------------------------------
class VectorIndex:
    def __init__(self, embedder=None):
        self.embedder = embedder or get_embedder()
        self.chunks: list[Chunk] = []
        self.M: np.ndarray | None = None

    def build(self, chunks: list[Chunk]):
        self.chunks = chunks
        texts = [c.text for c in chunks]
        if hasattr(self.embedder, "fitted"):
            self.embedder.fit(texts)
        self.M = self.embedder.encode(texts)
        return self

    def search(self, query: str, k: int = 5, where=None):
        q = self.embedder.encode([query])[0]
        sims = self.M @ q
        order = np.argsort(-sims)
        out = []
        for i in order:
            c = self.chunks[i]
            if where and not where(c):
                continue
            out.append((float(sims[i]), c))
            if len(out) == k:
                break
        return out

    def save(self, path: pathlib.Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        np.save(path.with_suffix(".npy"), self.M)
        path.with_suffix(".jsonl").write_text(
            "\n".join(json.dumps(c.to_json()) for c in self.chunks), encoding="utf-8")


def apply_ocr(docs: list[Doc], ocr_fn, raw_dir: pathlib.Path = RAW) -> list[Doc]:
    """The production answer to the Lab 1 silent failure: any PDF that extracted
    zero characters is sent to OCR instead of being indexed as nothing.
    ocr_fn(path) returns the recognised text."""
    by_source: dict[str, list[Doc]] = {}
    for d in docs:
        by_source.setdefault(d.meta.get("source", "?"), []).append(d)
    out = []
    for src, ds in by_source.items():
        empty = sum(len(d.text.strip()) for d in ds) == 0
        if empty and src.lower().endswith(".pdf"):
            text = ocr_fn(raw_dir / src)
            out.append(Doc(text=text, meta={"source": src, "page": 1, "format": "pdf", "ocr": True}))
        else:
            out.extend(ds)
    return out


def current_only(c: Chunk) -> bool:
    """The filter that would have prevented the Week 1 bank story."""
    return not c.meta.get("superseded", False)
